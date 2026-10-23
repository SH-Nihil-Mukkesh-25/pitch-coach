"""Unit and integration tests for Feature 8: Slide Coach Upgrades.

Verifies:
- Design critique fields (too_much_text, weak_title, missing_cta, visual_hierarchy_rating).
- Story-arc check detecting missing 'Ask' slide and flagging it.
- PPTX to PDF conversion and analysis pipeline parity.
"""
import io
import pytest
import pymupdf
import pptx
from pydantic import ValidationError

from slides.schemas import DeckSummary, SlideVisionDescription, SlideCoaching
from slides.pipeline import audit_deck_story_arc, convert_pptx_to_pdf


def _create_sample_pptx(slides_content: list[tuple[str, str]]) -> bytes:
    """Helper to generate an in-memory PPTX presentation."""
    prs = pptx.Presentation()
    # 16:9 aspect ratio standard
    prs.slide_width = pptx.util.Inches(10)
    prs.slide_height = pptx.util.Inches(5.625)

    for title_text, body_text in slides_content:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
        if slide.shapes.title:
            slide.shapes.title.text = title_text
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = body_text

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_to_pdf_conversion():
    """Verify PPTX binary is faithfully converted to PDF pages with extractable text."""
    deck_content = [
        ("The Problem in Healthcare", "Patients wait 45 minutes on average to see a physician."),
        ("Our Solution: MediQueue", "Instant AI triage reduces waiting times by 80%."),
        ("Traction & Growth", "12 enterprise hospital pilots signed, $40k MRR."),
    ]
    pptx_bytes = _create_sample_pptx(deck_content)
    assert len(pptx_bytes) > 0

    # Convert to PDF
    pdf_bytes = convert_pptx_to_pdf(pptx_bytes)
    assert pdf_bytes.startswith(b"%PDF")

    # Verify PyMuPDF can open and inspect text on each slide
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    assert len(doc) == 3

    page0_text = doc[0].get_text()
    assert "Healthcare" in page0_text or "Problem" in page0_text
    assert "Patients" in page0_text

    page1_text = doc[1].get_text()
    assert "MediQueue" in page1_text

    page2_text = doc[2].get_text()
    assert "Traction" in page2_text
    doc.close()


def test_story_arc_flags_missing_ask_slide():
    """
    Checkpoint 8 Core Verification:
    A deck with Problem, Solution, and Traction slides but MISSING an Ask slide
    must be flagged with has_ask = False, 'Ask' in missing_pillars and story_arc_warnings.
    """
    slides = [
        SlideVisionDescription(
            slide_number=1,
            title="The Urgent Problem",
            main_message="Current delivery infrastructure is broken and inefficient.",
            visible_text="Broken supply chain causes $10B annual losses.",
            visuals="Supply chain flow diagram",
            text_density="low",
            too_much_text=False,
            weak_title=False,
            missing_cta=False,
            visual_hierarchy_rating="good",
        ),
        SlideVisionDescription(
            slide_number=2,
            title="Our Solution: Autonomous Delivery",
            main_message="Fleet of autonomous drones delivering under 15 minutes.",
            visible_text="Drone hardware and software architecture.",
            visuals="Drone illustration and dispatch map",
            text_density="medium",
            too_much_text=False,
            weak_title=False,
            missing_cta=False,
            visual_hierarchy_rating="excellent",
        ),
        SlideVisionDescription(
            slide_number=3,
            title="Early Traction and User Growth",
            main_message="Over 5,000 deliveries completed across 3 pilot cities.",
            visible_text="Monthly delivery volume hockey stick curve.",
            visuals="Bar chart showing 300% MoM growth",
            text_density="low",
            too_much_text=False,
            weak_title=False,
            missing_cta=False,
            visual_hierarchy_rating="good",
        ),
    ]

    initial_summary = DeckSummary(
        story_arc="Strong problem and traction narrative.",
        missing_sections=[],
        overall_score=78,
        top_3_fixes=["Include a clear call to action"],
        has_problem=False,
        has_solution=False,
        has_traction=False,
        has_ask=False,
        missing_pillars=[],
        story_arc_warnings=[],
    )

    audited = audit_deck_story_arc(slides, initial_summary)

    # Core assertions
    assert audited.has_problem is True
    assert audited.has_solution is True
    assert audited.has_traction is True
    assert audited.has_ask is False, "Missing 'Ask' slide must be detected as False!"
    assert "Ask" in audited.missing_pillars, "'Ask' must be listed in missing narrative pillars!"
    assert any("Missing Ask Slide" in w for w in audited.story_arc_warnings)
    assert "Ask" in audited.missing_sections


def test_story_arc_complete_deck():
    """Verify that a deck with Problem, Solution, Traction, and Ask has no missing pillars."""
    slides = [
        SlideVisionDescription(
            slide_number=1,
            title="Problem",
            main_message="Inefficient scheduling.",
        ),
        SlideVisionDescription(
            slide_number=2,
            title="Our Solution",
            main_message="AI scheduling engine.",
        ),
        SlideVisionDescription(
            slide_number=3,
            title="Key Traction & Revenue",
            main_message="$50k MRR with 95% retention.",
        ),
        SlideVisionDescription(
            slide_number=4,
            title="The Ask: Raising $1.5M Seed Round",
            main_message="Seeking $1.5M to scale engineering and sales.",
            visible_text="Use of funds: 60% engineering, 40% GTM.",
        ),
    ]

    initial_summary = DeckSummary(
        story_arc="Complete and persuasive pitch arc.",
        missing_sections=[],
        overall_score=92,
        top_3_fixes=[],
    )

    audited = audit_deck_story_arc(slides, initial_summary)
    assert audited.has_problem is True
    assert audited.has_solution is True
    assert audited.has_traction is True
    assert audited.has_ask is True
    assert len(audited.missing_pillars) == 0


def test_slide_design_critique_schema():
    """Verify design critique fields and visual hierarchy ratings."""
    slide = SlideVisionDescription(
        slide_number=1,
        title="Overview",
        main_message="Company overview",
        visible_text="Bullet 1\nBullet 2\nBullet 3\nBullet 4\nBullet 5\nBullet 6\nBullet 7",
        visuals="None",
        text_density="high",
        readability_issues=["Small font size", "Low contrast"],
        too_much_text=True,
        weak_title=True,
        missing_cta=True,
        visual_hierarchy_rating="poor",
        design_critique=["Title 'Overview' is too generic", "Exceeds 6 bullet lines", "Lacks visual hierarchy"],
    )

    assert slide.too_much_text is True
    assert slide.weak_title is True
    assert slide.missing_cta is True
    assert slide.visual_hierarchy_rating == "poor"
    assert len(slide.design_critique) == 3

    # Rehydrated SlideCoaching
    coaching = SlideCoaching(
        slide_number=slide.slide_number,
        slide_title=slide.title,
        target_duration_seconds=30,
        key_takeaway="Company overview",
        talking_points=["Point 1", "Point 2"],
        transition_to_next="Next slide hook",
        slide_critique=["Reduce bullets"],
        too_much_text=slide.too_much_text,
        weak_title=slide.weak_title,
        missing_cta=slide.missing_cta,
        visual_hierarchy_rating=slide.visual_hierarchy_rating,
        design_critique=slide.design_critique,
    )
    assert coaching.too_much_text is True
    assert coaching.visual_hierarchy_rating == "poor"


def test_story_arc_hackathon_audience():
    """Verify that a hackathon technical deck is not penalized for missing funding ask or investor traction."""
    slides = [
        SlideVisionDescription(
            slide_number=1,
            title="The Problem: Heavy Cloud Datacenter Costs",
            main_message="Cloud servers are expensive and have high latency for local edge devices.",
        ),
        SlideVisionDescription(
            slide_number=2,
            title="System Architecture: STM32 + ESP8266",
            main_message="Hardware pipeline with USART2 at 115200 baud and SPI microSD.",
            visible_text="Bare-metal CMSIS drivers with zero HAL overhead.",
        ),
        SlideVisionDescription(
            slide_number=3,
            title="Working Prototype & Live Benchmarks",
            main_message="Tested 100% request throughput and under 128KB RAM footprint.",
        ),
        SlideVisionDescription(
            slide_number=4,
            title="Roadmap & Next Steps",
            main_message="Adding encrypted payload support and Bluetooth LE fallback.",
        ),
    ]

    summary = DeckSummary(
        story_arc="Strong technical hardware build.",
        missing_sections=[],
        overall_score=85,
        top_3_fixes=[],
        has_problem=False,
        has_solution=False,
        has_traction=False,
        has_ask=False,
        missing_pillars=[],
        story_arc_warnings=[],
    )

    audited = audit_deck_story_arc(slides, summary, audience="hackathon_judges")

    assert audited.has_problem is True
    assert audited.has_solution is True
    assert audited.has_traction is True  # Benchmarks / working prototype counted as traction
    assert audited.has_ask is True       # Roadmap / next steps counted as ask
    assert len(audited.missing_pillars) == 0
    # Crucially, must not contain generic investor funding ask warning
    assert not any("funding" in w.lower() for w in audited.story_arc_warnings)
    assert not any("Missing Ask Slide" in w for w in audited.story_arc_warnings)
